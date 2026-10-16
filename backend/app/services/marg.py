"""Read Marg ERP "STOCK & SALES ANALYSIS" Excel reports (.xls / .xlsx).

Layout of the report (one sheet):
  * page headers: shop name, address, phone, report title with the period,
    column headings (ITEM DESCRIPTION ... CLOSING STOCK M.EXP)
  * a company name on its own row, followed by that company's items
  * item rows: fixed-width description "NAME (20) PACK (7) UNIT", then numbers;
    the last column is "closing quantity  M/YY" (earliest expiry) or 0
  * " Quantity" / " Value in Rs." total rows, page breaks ("Continued ..",
    "Page No..") and a "PURCHASE DETAIL" block, which are all skipped

The report has no batch numbers and no prices, so none are created.
"""
from __future__ import annotations

import calendar
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

CLOSING_RE = re.compile(r"^\s*(-?\d+(?:\.\d+)?)\s+(\d{1,2})/(\d{2,4})\s*$")
PERIOD_RE = re.compile(r"(\d{2})-(\d{2})-(\d{4})\s*-\s*(\d{2})-(\d{2})-(\d{4})")
SKIP_PREFIXES = ("STOCK & SALES ANALYSIS", "ITEM DESCRIPTION", "PURCHASE DETAIL", "SUPPLIER NAME",
                 "I AM SATISFIED WITH MARG", "PHONE :", "PAGE NO", "CONTINUED")


@dataclass
class MargItem:
    row: int
    company: str
    name: str
    pack: str
    unit: str
    closing: int
    earliest_expiry: date | None


@dataclass
class MargReport:
    shop_name: str = ""
    period_from: date | None = None
    period_to: date | None = None
    items: list[MargItem] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _cells(content: bytes, filename: str) -> list[list]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        return [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheet = max(book.sheets(), key=lambda s: s.nrows)  # the report sheet, not the empty ones
    return [sheet.row_values(r) for r in range(sheet.nrows)]


def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _end_of_month(month: int, year: int) -> date:
    if year < 100:
        year += 2000
    return date(year, month, calendar.monthrange(year, month)[1])


def _split_description(text: str) -> tuple[str, str, str]:
    """'VEINFLUX FILL PARTI 5ML    PCS' -> ('VEINFLUX FILL PARTI', '5ML', 'PCS').

    Normal rows are fixed width (name 20, pack 7, unit). When a long name pushes the
    columns over, the parts are separated by two or more spaces instead.
    """
    raw = text.rstrip()
    if len(raw) == 30 and raw.endswith("PCS") and raw[19] == " " and raw[26] == " ":
        name, pack, unit = raw[:20], raw[20:27], raw[27:]
        return " ".join(name.split()), " ".join(pack.split()), unit.strip()
    parts = [p for p in re.split(r"\s{2,}", raw.strip()) if p]
    if not parts:
        return "", "", ""
    unit = ""
    if len(parts) > 1 and parts[-1] == "PCS":
        unit = parts.pop()
    name = " ".join(parts[0].split())
    pack = " ".join(parts[1].split()) if len(parts) > 1 else ""
    return name, pack, unit


def parse_marg_report(content: bytes, filename: str = "report.xls") -> MargReport:
    rows = _cells(content, filename)
    rep = MargReport()
    if rows and rows[0] and isinstance(rows[0][0], str):
        rep.shop_name = rows[0][0].strip()

    company: str | None = None
    in_purchase_block = False
    for r, cells in enumerate(rows, start=1):
        cells = list(cells) + [""] * (12 - len(cells))
        first = cells[0] if isinstance(cells[0], str) else str(cells[0] or "")
        upper = first.strip().upper()
        rest = [c for c in cells[1:] if c not in ("", None)]

        if upper.startswith("STOCK & SALES ANALYSIS") and rep.period_to is None:
            m = PERIOD_RE.search(upper)
            if m:
                d1, m1, y1, d2, m2, y2 = map(int, m.groups())
                rep.period_from, rep.period_to = date(y1, m1, d1), date(y2, m2, d2)
        if upper.startswith("PURCHASE DETAIL"):
            in_purchase_block = True
            continue
        if in_purchase_block:
            if upper.startswith("I AM SATISFIED WITH MARG") or (not rest and upper and not upper.startswith(SKIP_PREFIXES)):
                in_purchase_block = False
            else:
                continue
        if not upper or first.startswith(" ") or upper.startswith(SKIP_PREFIXES) or upper == rep.shop_name.upper():
            continue
        if r <= 4 and not rest:
            continue  # address line in the page header

        if not rest:
            company = " ".join(first.split())
            continue

        closing_cell = cells[11]
        if not _is_number(cells[1]):
            continue  # not an item row
        if company is None:
            rep.warnings.append(f"Row {r}: item before any company name, skipped")
            continue
        name, pack, unit = _split_description(first)
        if not name:
            continue
        expiry = None
        if _is_number(closing_cell) or re.fullmatch(r"\s*-?\d+(\.\d+)?\s*", str(closing_cell)):
            qty = float(closing_cell)
        else:
            m = CLOSING_RE.match(str(closing_cell))
            if not m:
                if str(closing_cell).strip():
                    rep.warnings.append(f"Row {r}: could not read closing stock '{closing_cell}' for {name}, skipped")
                continue
            qty = float(m.group(1))
            month, year = int(m.group(2)), int(m.group(3))
            if 1 <= month <= 12:
                expiry = _end_of_month(month, year)
            else:
                rep.warnings.append(f"Row {r}: bad expiry '{closing_cell}' for {name}")
        if qty < 0:
            rep.warnings.append(f"Row {r}: negative closing stock ({qty:g}) for {name}; shown as 0")
            qty = 0
        if qty != int(qty):
            rep.warnings.append(f"Row {r}: fractional stock {qty:g} for {name}; rounded down")
        rep.items.append(MargItem(r, company, name, pack, unit, int(qty), expiry))
    return rep


def period_label(rep: MargReport) -> str:
    if rep.period_to:
        return f"Marg stock report up to {rep.period_to:%d-%m-%Y}"
    return f"Marg stock report imported {datetime.now():%d-%m-%Y}"
